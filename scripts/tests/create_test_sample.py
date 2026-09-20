"""
Creates a realistic sample Word docx representing lecture notes exported from Google Docs,
including Yellow highlights, Green definitions, and Blue notes.
"""
import docx
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

doc = docx.Document()
doc.add_heading("PSYC 3590: Drugs & Behaviour - Lecture 2: Pharmacodynamics & Addiction", level=1)

# Paragraph 1: Yellow term + Green definition
p1 = doc.add_paragraph()
r1_1 = p1.add_run("Tolerance")
r1_1._r.get_or_add_rPr().append(parse_xml(f'<w:highlight {nsdecls("w")} w:val="yellow"/>'))
p1.add_run(" is defined as ")
r1_2 = p1.add_run("a state of progressively decreasing responsiveness to a drug following repeated administration, requiring higher doses to achieve the same effect.")
r1_2._r.get_or_add_rPr().append(parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="B7E1CD"/>'))

# Paragraph 2: Standalone Yellow (Important mechanism)
p2 = doc.add_paragraph()
p2.add_run("Key principle of antipsychotic efficacy: ")
r2_1 = p2.add_run("Dopamine D2 receptor occupancy between 65% and 80% is required for clinical therapeutic response, whereas occupancy exceeding 80% sharply increases the risk of extrapyramidal symptoms (EPS).")
r2_1._r.get_or_add_rPr().append(parse_xml(f'<w:highlight {nsdecls("w")} w:val="yellow"/>'))

# Paragraph 3: Standalone Green definition with plain text term
p3 = doc.add_paragraph()
p3.add_run("The therapeutic index (TI) of a drug represents ")
r3_1 = p3.add_run("the ratio between the toxic dose (TD50) or lethal dose (LD50) and the therapeutically effective dose (ED50), serving as a quantitative measurement of relative safety.")
r3_1._r.get_or_add_rPr().append(parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="00FF00"/>'))

# Paragraph 4: Yellow finding with Blue clinical note
p4 = doc.add_paragraph()
p4.add_run("Receptor downregulation: ")
r4_1 = p4.add_run("Chronic agonist stimulation induces endocytosis and degradation of cell surface receptors.")
r4_1._r.get_or_add_rPr().append(parse_xml(f'<w:highlight {nsdecls("w")} w:val="yellow"/>'))
r4_2 = p4.add_run(" (e.g. chronic opioid exposure leading to cellular tolerance at mu-opioid receptors)")
r4_2._r.get_or_add_rPr().append(parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="C9DAF8"/>'))

doc.save("sample_lecture_notes.docx")
print("Saved sample_lecture_notes.docx")
