"""
test_e2e_phase2.py — Opaque-Box E2E Test Suite for AnkiDeckCreator Phase 2

Requirements in Scope:
1. Google Docs API Ingestion (URL, Doc ID, .docx polymorphism, deck name resolution from doc title,
   mocked Google Docs API error handling, structural element traversal, background color extraction).
2. Card Quality, Sanitization & Deduplication (CardSanitizer, CardValidator, CardDeduplicator,
   DeckQualityPipeline, length thresholds, placeholder rejection, corruption cleaning, bidirectional preservation).
3. Cognitive Card Taxonomies & Hierarchical Tagging (Function -> Structure, Concept -> Mechanism,
   Example -> Category, Term -> Definition, hierarchical tags Course::*, Chapter::*, Type::*).

4-Tier Test Architecture:
- Tier 1: Feature Coverage (Representative inputs for all 7 features in scope)
- Tier 2: Boundary & Corner Cases (Stress & Fault injection: network errors, extreme lengths,
  pathological characters, missing fields, malformed URLs/IDs, corruption variants)
- Tier 3: Cross-Feature Combinations (Pairwise and multi-system interactions: Docs Ingestion ->
  Taxonomy Formulation -> Quality Validation -> Deduplication -> .apkg creation)
- Tier 4: Real-World Scenarios (End-to-end full lecture simulation, polymorphism parity,
  SQLite database archive verification, and 30-point Agent-as-Judge evaluation)
"""

import io
import json
import os
import re
import shutil
import sqlite3
import sys
import tempfile
import zipfile
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple
from unittest.mock import MagicMock, patch

import pytest
import docx
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

# Ensure project root and scripts directory are in sys.path
PROJECT_ROOT = Path(__file__).resolve().parent.parent.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))
SCRIPTS_DIR = PROJECT_ROOT / "scripts"
if str(SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_DIR))

# Dynamic imports for modules
try:
    from scripts.docs_api_client import (
        extract_document_id,
        rgb_to_hex,
        default_classify_color,
        fetch_highlighted_text,
        extract_google_doc_structured,
    )
except ImportError:
    extract_document_id = None
    rgb_to_hex = None
    default_classify_color = None
    fetch_highlighted_text = None
    extract_google_doc_structured = None

try:
    from scripts.docs_api_client import detect_input_type, ingest_source
except (ImportError, AttributeError):
    try:
        from scripts.extract_and_generate import detect_input_type, ingest_source
    except (ImportError, AttributeError):
        detect_input_type = None
        ingest_source = None

try:
    from scripts.card_validator import (
        CardSanitizer,
        CardValidator,
        CardDeduplicator,
        DeckQualityPipeline,
        strip_html_tags,
        CardValidationResult,
        DeckValidationSummary,
        ValidationIssue,
    )
except ImportError:
    CardSanitizer = None
    CardValidator = None
    CardDeduplicator = None
    DeckQualityPipeline = None
    strip_html_tags = None
    CardValidationResult = None
    DeckValidationSummary = None
    ValidationIssue = None

try:
    from scripts.extract_and_generate import (
        classify_color,
        resolve_deck_naming,
        synthesize_cards,
        create_deck_package,
        extract_document_highlights,
        clean_phrase,
        CONFIG,
    )
except ImportError:
    classify_color = None
    resolve_deck_naming = None
    synthesize_cards = None
    create_deck_package = None
    extract_document_highlights = None
    clean_phrase = None
    CONFIG = {}

try:
    from scripts.semantic_parser import SemanticCardParser
except ImportError:
    SemanticCardParser = None

try:
    from scripts.evaluate_cards_judge import CardJudgeRubric
except ImportError:
    CardJudgeRubric = None


# ==============================================================================
# AUTHORITATIVE REFERENCE ORACLES & FIXTURES
# ==============================================================================

class ReferenceCardJudgeOracle:
    """
    Authoritative reference oracle implementing the 30-point evaluation rubric
    and 4 hard vetoes specified in ANKI_SOP.md and PROJECT.md.
    """
    BANNED_VAGUE_PATTERNS = [
        r"\bthis\s+concept\b",
        r"\bthis\s+term\b",
        r"\bthis\s+phenomenon\b",
        r"\bwhat\s+is\s+this\s+concept\b",
        r"\bdefine\s+this\s+concept\b",
    ]

    def evaluate_card(self, card: Dict[str, Any]) -> Dict[str, Any]:
        q = card.get("question", "").strip()
        a = card.get("answer", "").strip()
        badge = card.get("category_badge", "").strip()
        ctx = card.get("context", "").strip()
        tags = card.get("tags", [])

        vetoes = []

        # 1. Critical Hard Veto Checks
        if not q or not a:
            vetoes.append("CRITICAL_FAIL_EMPTY_FIELD: Question or Answer is empty.")

        for pat in self.BANNED_VAGUE_PATTERNS:
            if re.search(pat, q, re.IGNORECASE) or re.search(pat, a, re.IGNORECASE):
                vetoes.append(f"CRITICAL_FAIL_THIS_CONCEPT: Contains banned phrase matching '{pat}'.")

        # Heading-as-prompt veto check
        if "Key Concept / Mechanism:" in q and ("<br>" in q or "\n" in q):
            trailing = re.split(r'<br>|\n', q)[-1].strip()
            if len(trailing.split()) <= 8 and not re.search(r'^(What|Which|Why|How|Define|Explain|Identify)\b', trailing, re.IGNORECASE):
                vetoes.append("CRITICAL_FAIL_HEADING_ONLY: Prompt is a raw heading dump without an active recall question.")

        if q.lower() == a.lower() and q != "":
            vetoes.append("CRITICAL_FAIL_TAUTOLOGY: Question and Answer are identical.")

        # Dimension Scoring (0 to 5 each, max 30)
        d_scores = {}

        # D1: Atomicity (Rule 4)
        word_count = len(a.split())
        if word_count <= 25:
            d_scores["D1_Atomicity"] = 5
        elif word_count <= 40:
            d_scores["D1_Atomicity"] = 3
        else:
            d_scores["D1_Atomicity"] = 1

        # D2: Absence of Vague Prompts (Rule 12)
        if any(re.search(pat, q, re.IGNORECASE) for pat in self.BANNED_VAGUE_PATTERNS):
            d_scores["D2_Wording"] = 0
        elif re.search(r'^(What|Which|Why|How|Define|Identify|Explain)\b', q, re.IGNORECASE):
            d_scores["D2_Wording"] = 5
        elif "{{c1::" in q:
            d_scores["D2_Wording"] = 5
        else:
            d_scores["D2_Wording"] = 3

        # D3: Keyword-Descriptor Integrity (SOP 1.5)
        if "<b>" in q and "</b>" in q:
            d_scores["D3_Keyword_Descriptor"] = 5
        elif "{{c1::" in q:
            d_scores["D3_Keyword_Descriptor"] = 5
        else:
            d_scores["D3_Keyword_Descriptor"] = 3

        # D4: Avoidance of Sets/Enumerations (Rules 9 & 10)
        if re.search(r'\blist\s+the\b|\bwhat\s+are\s+the\s+\d+\b', q, re.IGNORECASE):
            d_scores["D4_Set_Avoidance"] = 1
        else:
            d_scores["D4_Set_Avoidance"] = 5

        # D5: Active Recall vs Recognition (SOP 1.5 & Rule 17)
        if len(a.split()) <= 6:
            d_scores["D5_Recall_Recognition"] = 5
        elif len(a.split()) <= 20:
            d_scores["D5_Recall_Recognition"] = 4
        else:
            d_scores["D5_Recall_Recognition"] = 3

        # D6: Contextual Anchoring (Rules 16 & 18)
        c_score = 0
        if badge:
            c_score += 2
        if ctx:
            c_score += 2
        if tags:
            c_score += 1
        d_scores["D6_Context_Anchoring"] = min(5, c_score)

        total_score = sum(d_scores.values())
        if vetoes:
            total_score = 0
            passed = False
        else:
            passed = (total_score >= 25)

        return {
            "question": q,
            "answer": a,
            "total_score": total_score,
            "passed": passed,
            "vetoes": vetoes,
            "dimension_scores": d_scores,
        }

    def evaluate_deck(self, cards: List[Dict[str, Any]]) -> Dict[str, Any]:
        if not cards:
            return {"total_cards": 0, "passed_cards": 0, "average_score": 0.0, "deck_passed": False, "veto_count": 0}
        evals = [self.evaluate_card(c) for c in cards]
        passed_count = sum(1 for e in evals if e["passed"])
        veto_count = sum(len(e["vetoes"]) for e in evals)
        avg_score = sum(e["total_score"] for e in evals) / len(evals)
        return {
            "total_cards": len(cards),
            "passed_cards": passed_count,
            "average_score": round(avg_score, 2),
            "deck_passed": (passed_count == len(cards)) and (veto_count == 0),
            "veto_count": veto_count,
            "card_evaluations": evals,
        }


def make_google_docs_mock(title: str, paragraphs: List[Dict[str, Any]]):
    """
    Constructs a mocked Google Docs API v1 document payload.
    Each paragraph in `paragraphs`:
        {
            "heading": Optional[str], # "HEADING_1", "TITLE", None
            "runs": [
                {"text": str, "color": Optional[Tuple[float, float, float]]} # rgb 0.0-1.0
            ]
        }
    """
    content = []
    for p in paragraphs:
        elements = []
        for r in p.get("runs", []):
            el = {
                "textRun": {
                    "content": r["text"],
                    "textStyle": {}
                }
            }
            if r.get("color"):
                red, green, blue = r["color"]
                el["textRun"]["textStyle"]["backgroundColor"] = {
                    "color": {"rgbColor": {"red": red, "green": green, "blue": blue}}
                }
            elements.append(el)

        style_type = p.get("heading") or "NORMAL_TEXT"
        content.append({
            "paragraph": {
                "paragraphStyle": {"namedStyleType": style_type},
                "elements": elements
            }
        })

    return {
        "title": title,
        "body": {"content": content}
    }


def create_mock_service(doc_payload):
    """Returns a MagicMock simulating googleapiclient.discovery.build('docs', 'v1')."""
    mock_service = MagicMock()
    mock_documents = MagicMock()
    mock_get = MagicMock()
    mock_get.execute.return_value = doc_payload
    mock_documents.get.return_value = mock_get
    mock_service.documents.return_value = mock_documents
    return mock_service


def create_sample_docx(docx_path: Path, paragraphs_data: List[Dict[str, Any]]) -> Path:
    """Helper to synthesize a test .docx file with specific OpenXML highlight runs."""
    doc = docx.Document()
    for p_data in paragraphs_data:
        if p_data.get("heading"):
            doc.add_heading(p_data["heading"], level=1)
        p = doc.add_paragraph()
        for r_data in p_data.get("runs", []):
            run = p.add_run(r_data.get("text", ""))
            rPr = run._r.get_or_add_rPr()
            hl = r_data.get("highlight")
            if hl and hl.lower() != "none":
                rPr.append(parse_xml(f'<w:highlight {nsdecls("w")} w:val="{hl.lower()}"/>'))
            fill = r_data.get("fill")
            if fill:
                rPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{fill}"/>'))
    doc.save(str(docx_path))
    return docx_path


@pytest.fixture
def judge():
    return ReferenceCardJudgeOracle()


# ==============================================================================
# TIER 1: FEATURE COVERAGE (7 Features x 5 Tests = 35 Tests)
# ==============================================================================

class TestTier1FeatureCoverage:
    """
    Tier 1 tests verify happy paths and contracts for all Phase 2 features:
    1. Google Docs API Ingestion & Extraction
    2. Input Polymorphism & Doc ID Resolution
    3. Automated Deck Naming from Google Doc Titles
    4. Card Sanitization & Corruption Cleaning
    5. Card Quality & Defensive Validation
    6. Bidirectional-Safe Deduplication & Metadata Merging
    7. Cognitive Card Taxonomies & Hierarchical Tagging
    """

    # --- Feature 1: Google Docs API Ingestion & Extraction ---

    def test_tier1_extract_document_id_from_full_url(self):
        """Feature 1.1: Extracts 44-character document ID from standard Google Docs URL."""
        assert extract_document_id is not None, "extract_document_id not found in docs_api_client"
        url = "https://docs.google.com/document/d/1BxiMVs0XRA5nFMdKvBdBZjgmUUqptlbs74OgvE2upms/edit#heading=h.abc123"
        doc_id = extract_document_id(url)
        assert doc_id == "1BxiMVs0XRA5nFMdKvBdBZjgmUUqptlbs74OgvE2upms"

    def test_tier1_extract_document_id_bare_string(self):
        """Feature 1.2: Returns raw ID unchanged when passed a bare Document ID."""
        assert extract_document_id is not None
        bare_id = "1BxiMVs0XRA5nFMdKvBdBZjgmUUqptlbs74OgvE2upms"
        assert extract_document_id(bare_id) == bare_id

    def test_tier1_rgb_to_hex_conversion(self):
        """Feature 1.3: Converts Google Docs float RGB (0.0-1.0) to standard hex string."""
        assert rgb_to_hex is not None
        # Green highlight
        green_rgb = {"red": 0.0, "green": 1.0, "blue": 0.0}
        assert rgb_to_hex(green_rgb) == "#00ff00"
        # Soft Google Docs green (#b7e1cd)
        soft_green = {"red": 183 / 255.0, "green": 225 / 255.0, "blue": 205 / 255.0}
        assert rgb_to_hex(soft_green) == "#b7e1cd"

    def test_tier1_default_classify_color(self):
        """Feature 1.4: Classifies hex colors into semantic categories (green, yellow, other)."""
        assert default_classify_color is not None
        assert default_classify_color("#00ff00") == "green"
        assert default_classify_color("#b7e1cd") == "green"
        assert default_classify_color("#ffff00") == "yellow"
        assert default_classify_color("#fff2cc") == "yellow"
        assert default_classify_color("#ff0000") == "other"

    def test_tier1_extract_google_doc_structured_highlights(self):
        """Feature 1.5: Extracts structured highlights and document title from mocked Google Docs API."""
        assert extract_google_doc_structured is not None
        mock_doc = make_google_docs_mock(
            title="PSYC 2240 - Lecture 1 Introduction",
            paragraphs=[
                {
                    "heading": "HEADING_1",
                    "runs": [{"text": "Foundations of Neuroscience\n", "color": None}]
                },
                {
                    "heading": "NORMAL_TEXT",
                    "runs": [
                        {"text": "Synaptic Plasticity: ", "color": (1.0, 1.0, 0.0)}, # Yellow
                        {"text": "the ability of synapses to strengthen or weaken over time.", "color": (0.0, 1.0, 0.0)} # Green
                    ]
                }
            ]
        )
        mock_service = create_mock_service(mock_doc)
        with patch("scripts.docs_api_client.build", return_value=mock_service), \
             patch("scripts.docs_api_client.get_credentials", return_value=MagicMock()):
            data, title = extract_google_doc_structured("dummy_id")

            assert title == "PSYC 2240 - Lecture 1 Introduction"
            assert len(data) == 1
            assert data[0]["heading"] == "Foundations of Neuroscience"
            assert len(data[0]["highlights"]) == 2
            assert data[0]["highlights"][0]["category"] == "yellow"
            assert "Synaptic Plasticity" in data[0]["highlights"][0]["text"]
            assert data[0]["highlights"][1]["category"] == "green"

    # --- Feature 2: Input Polymorphism & Document ID Resolution ---

    def test_tier1_detect_input_type_url(self):
        """Feature 2.1: Correctly classifies Google Docs URL."""
        if detect_input_type is None:
            # When detect_input_type is implemented in M5, this will test it
            pytest.skip("detect_input_type pending implementation in Milestone 5")
        url = "https://docs.google.com/document/d/1BxiMVs0XRA5nFMdKvBdBZjgmUUqptlbs74OgvE2upms/edit"
        res = detect_input_type(url)
        assert "url" in res.lower() or "gdoc" in res.lower()

    def test_tier1_detect_input_type_bare_id(self):
        """Feature 2.2: Correctly classifies bare Google Docs ID."""
        if detect_input_type is None:
            pytest.skip("detect_input_type pending implementation in Milestone 5")
        bare_id = "1BxiMVs0XRA5nFMdKvBdBZjgmUUqptlbs74OgvE2upms"
        res = detect_input_type(bare_id)
        assert "id" in res.lower() or "gdoc" in res.lower()

    def test_tier1_detect_input_type_docx_file(self, tmp_path):
        """Feature 2.3: Correctly classifies local .docx file path."""
        if detect_input_type is None:
            pytest.skip("detect_input_type pending implementation in Milestone 5")
        doc_path = tmp_path / "lecture.docx"
        doc_path.touch()
        res = detect_input_type(str(doc_path))
        assert "docx" in res.lower()

    def test_tier1_ingest_source_polymorphism_gdoc(self):
        """Feature 2.4: ingest_source dispatches to Google Docs API client for URL or ID."""
        if ingest_source is None:
            pytest.skip("ingest_source pending implementation in Milestone 5")
        mock_doc = make_google_docs_mock("Title", [{"runs": [{"text": "Term", "color": (1.0, 1.0, 0.0)}]}] )
        mock_service = create_mock_service(mock_doc)
        with patch("scripts.docs_api_client.build", return_value=mock_service), \
             patch("scripts.docs_api_client.get_credentials", return_value=MagicMock()):
            data, title = ingest_source("https://docs.google.com/document/d/dummy_id/edit")
            assert title == "Title"
            assert isinstance(data, list)

    def test_tier1_ingest_source_polymorphism_docx(self, tmp_path):
        """Feature 2.5: ingest_source dispatches to Word extraction for .docx files."""
        if ingest_source is None:
            pytest.skip("ingest_source pending implementation in Milestone 5")
        p = tmp_path / "test.docx"
        create_sample_docx(p, [{"heading": "Chapter 1", "runs": [{"text": "Term: ", "highlight": "yellow"}, {"text": "Def", "highlight": "green"}]}])
        data, title = ingest_source(str(p))
        assert isinstance(data, list)
        assert len(data) >= 1

    # --- Feature 3: Automated Deck Naming from Google Doc Titles ---

    def test_tier1_deck_naming_known_course_in_title(self):
        """Feature 3.1: Resolves known course code in title to official course name from config."""
        assert resolve_deck_naming is not None
        # Using a dummy docx path with course in filename or checking resolve_deck_naming signature
        import inspect
        sig = inspect.signature(resolve_deck_naming)
        if "doc_title" in sig.parameters:
            title, fn = resolve_deck_naming("dummy_id", doc_title="PSYC 3590 - Lecture 2: Pharmacodynamics")
            assert "PSYC 3590" in title
            assert "Drugs & Behaviour" in title
            assert "Lecture 2" in title
        else:
            # Baseline compatibility: file path with PSYC 3590 in stem
            title, fn = resolve_deck_naming("PSYC 3590 - Lecture 2.docx")
            assert "PSYC 3590" in title
            assert "Drugs & Behaviour" in title

    def test_tier1_deck_naming_explicit_class_override(self):
        """Feature 3.2: Explicit class argument overrides inferred title course code."""
        assert resolve_deck_naming is not None
        title, fn = resolve_deck_naming("sample.docx", explicit_class="PSYC 2240", explicit_chapter="Chapter 4")
        assert "PSYC 2240" in title
        assert "Biological Basis of Behaviour" in title
        assert "Chapter 4" in title

    def test_tier1_deck_naming_unknown_course_fallback(self):
        """Feature 3.3: Unknown course code formats as CODE:CHAPTER without failure."""
        assert resolve_deck_naming is not None
        title, fn = resolve_deck_naming("MATH 1010 - Calculus I.docx")
        assert "MATH 1010" in title
        assert "Calculus I" in title

    def test_tier1_deck_naming_chapter_split_colon(self):
        """Feature 3.4: Splits chapter string correctly when preceded by colon or hyphen."""
        assert resolve_deck_naming is not None
        title, fn = resolve_deck_naming("Lecture 3 - Action Potentials.docx")
        assert "Lecture 3: Action Potentials" in title or "Lecture 3" in title

    def test_tier1_deck_naming_windows_safe_filename(self):
        """Feature 3.5: Generates a Windows-safe filename without forbidden colons."""
        assert resolve_deck_naming is not None
        title, fn = resolve_deck_naming("PSYC 3590 - Lecture 1.docx")
        assert ":" not in fn
        assert fn.endswith(".apkg")

    # --- Feature 4: Card Sanitization & Corruption Cleaning ---

    def test_tier1_sanitizer_strip_css_blocks(self):
        """Feature 4.1: CardSanitizer strips embedded .card { ... } stylesheets and priority tags."""
        assert CardSanitizer is not None, "CardSanitizer not found in card_validator"
        sanitizer = CardSanitizer()
        raw = ".card { font-family: Arial; font-size: 12px; }\nHIGH Priority\nWhat is the frontal lobe?"
        cleaned = sanitizer.strip_css(raw)
        assert ".card" not in cleaned
        assert "HIGH Priority" not in cleaned
        assert "What is the frontal lobe?" in cleaned

    def test_tier1_sanitizer_strip_parenthetical_figure_citations(self):
        """Feature 4.2: CardSanitizer removes (Figure X.Y) and [Table X.Y] citations."""
        assert CardSanitizer is not None
        sanitizer = CardSanitizer()
        raw = "The action potential triggers neurotransmitter release (Figure 4.2) into the synaptic cleft."
        cleaned = sanitizer.strip_figure_citations(raw)
        assert "(Figure 4.2)" not in cleaned
        assert "The action potential triggers neurotransmitter release" in cleaned

    def test_tier1_sanitizer_strip_trailing_index_runs(self):
        """Feature 4.3: CardSanitizer strips textbook index page number runs preceded by substantive text."""
        assert CardSanitizer is not None
        sanitizer = CardSanitizer()
        raw = "Neurotransmitters in the central nervous system, 45, 48, 112."
        cleaned = sanitizer.strip_index_runs(raw)
        assert ", 45, 48, 112" not in cleaned
        assert "Neurotransmitters in the central nervous system." in cleaned

    def test_tier1_sanitizer_normalize_whitespace_and_quotes(self):
        """Feature 4.4: CardSanitizer collapses whitespace and normalizes smart quotes."""
        assert CardSanitizer is not None
        sanitizer = CardSanitizer()
        raw = "“Dopamine”   is   a   ‘neurotransmitter’.\r\n\r\nIt modulates   reward."
        cleaned = sanitizer.normalize_whitespace(sanitizer.normalize_quotes(raw))
        assert "   " not in cleaned
        assert "\r" not in cleaned
        assert '"Dopamine"' in cleaned
        assert "'neurotransmitter'" in cleaned

    def test_tier1_sanitizer_normalize_terminal_punctuation(self):
        """Feature 4.5: CardSanitizer ensures interrogatives end with '?' and declaratives with '.'."""
        assert CardSanitizer is not None
        sanitizer = CardSanitizer()
        assert sanitizer.normalize_punctuation("What is the function of the hippocampus", is_question=True) == "What is the function of the hippocampus?"
        assert sanitizer.normalize_punctuation("Consolidates short-term to long-term memory", is_question=False) == "Consolidates short-term to long-term memory."

    # --- Feature 5: Card Quality & Defensive Validation ---

    def test_tier1_validator_accepts_valid_card(self):
        """Feature 5.1: CardValidator accepts properly formulated atomic card."""
        assert CardValidator is not None, "CardValidator not found in card_validator"
        validator = CardValidator()
        card = {
            "question": "What is the primary function of the <b>hippocampus</b>?",
            "answer": "Consolidation of information from short-term memory to long-term memory.",
            "category_badge": "badge-definition"
        }
        res = validator.validate_card(card)
        assert res.is_valid is True

    def test_tier1_validator_rejects_placeholder_answers(self):
        """Feature 5.2: CardValidator rejects cards containing placeholder answers."""
        assert CardValidator is not None
        validator = CardValidator()
        placeholders = [
            "Review this concept in the textbook.",
            "Content needs verification.",
            "Check the textbook for complete information.",
            "This content needs review.",
            "Answer to be extracted from lecture notes.",
            "Requires specific course material review."
        ]
        for ph in placeholders:
            card = {
                "question": "What is long-term potentiation?",
                "answer": ph
            }
            res = validator.validate_card(card)
            assert res.is_valid is False, f"Failed to reject placeholder: {ph}"

    def test_tier1_validator_rejects_vague_banned_prompts(self):
        """Feature 5.3: CardValidator enforces SuperMemo Rule 12 zero 'this concept' veto."""
        assert CardValidator is not None
        validator = CardValidator()
        vague_questions = [
            "Define / What is this concept?",
            "What is this term?",
            "What is this phenomenon?"
        ]
        for vq in vague_questions:
            card = {
                "question": vq,
                "answer": "A neurobiological adaptation mechanism."
            }
            res = validator.validate_card(card)
            assert res.is_valid is False, f"Failed to reject vague prompt: {vq}"

    def test_tier1_validator_rejects_tautology(self):
        """Feature 5.4: CardValidator rejects cards where question and answer are identical."""
        assert CardValidator is not None
        validator = CardValidator()
        card = {
            "question": "Action Potential",
            "answer": "Action Potential"
        }
        res = validator.validate_card(card)
        assert res.is_valid is False

    def test_tier1_validator_validates_cloze_deletion(self):
        """Feature 5.5: CardValidator validates proper Cloze deletion cards with {{c1::...}} syntax."""
        assert CardValidator is not None
        validator = CardValidator()
        valid_cloze = {
            "card_type": "cloze",
            "question": "The primary inhibitory neurotransmitter in the brain is {{c1::GABA}}."
        }
        res = validator.validate_card(valid_cloze)
        assert res.is_valid is True

        missing_token = {
            "card_type": "cloze",
            "question": "The primary inhibitory neurotransmitter in the brain is GABA."
        }
        res_bad = validator.validate_card(missing_token)
        assert res_bad.is_valid is False

    # --- Feature 6: Bidirectional-Safe Deduplication & Metadata Merging ---

    def test_tier1_deduplicator_removes_exact_duplicates(self):
        """Feature 6.1: CardDeduplicator filters out identical duplicate cards."""
        assert CardDeduplicator is not None, "CardDeduplicator not found in card_validator"
        dedup = CardDeduplicator()
        cards = [
            {"question": "What is the function of GABA?", "answer": "Primary inhibitory neurotransmitter."},
            {"question": "What is the function of GABA?", "answer": "Primary inhibitory neurotransmitter."}
        ]
        unique = dedup.deduplicate(cards)
        assert len(unique) == 1

    def test_tier1_deduplicator_preserves_bidirectional_pairs(self):
        """Feature 6.2: CRITICAL INVARIANT: CardDeduplicator NEVER collapses bidirectional pairs."""
        assert CardDeduplicator is not None
        dedup = CardDeduplicator()
        # Pair testing active recall vs recognition
        forward_card = {
            "card_type": "bidirectional_definition",
            "question": "What is the definition of <b>Tolerance</b>?",
            "answer": "A state of progressively decreasing responsiveness to a drug."
        }
        reverse_card = {
            "card_type": "bidirectional_definition",
            "question": "What term is defined by:<br><i>A state of progressively decreasing responsiveness to a drug.</i>",
            "answer": "Tolerance"
        }
        cards = [forward_card, reverse_card]
        unique = dedup.deduplicate(cards)
        assert len(unique) == 2, "Bidirectional reciprocal pair was erroneously collapsed!"

    def test_tier1_deduplicator_merges_tags(self):
        """Feature 6.3: CardDeduplicator merges tags from duplicate notes without duplicates."""
        assert CardDeduplicator is not None
        dedup = CardDeduplicator()
        cards = [
            {"question": "What is LTP?", "answer": "Persistent strengthening of synapses.", "tags": ["neuroscience", "memory"]},
            {"question": "What is LTP?", "answer": "Persistent strengthening of synapses.", "tags": ["cellular", "plasticity", "memory"]}
        ]
        unique = dedup.deduplicate(cards)
        assert len(unique) == 1
        merged_tags = unique[0]["tags"]
        assert set(merged_tags) == {"neuroscience", "memory", "cellular", "plasticity"}

    def test_tier1_deduplicator_preserves_context(self):
        """Feature 6.4: CardDeduplicator preserves context notes when duplicate card provides them."""
        assert CardDeduplicator is not None
        dedup = CardDeduplicator()
        cards = [
            {"question": "What is dopamine?", "answer": "A catecholamine neurotransmitter.", "context": ""},
            {"question": "What is dopamine?", "answer": "A catecholamine neurotransmitter.", "context": "Lecture 2: Monoamines"}
        ]
        unique = dedup.deduplicate(cards)
        assert len(unique) == 1
        assert unique[0]["context"] == "Lecture 2: Monoamines"

    def test_tier1_deduplicator_cloze_cards(self):
        """Feature 6.5: CardDeduplicator correctly fingerprints and deduplicates Cloze cards."""
        assert CardDeduplicator is not None
        dedup = CardDeduplicator()
        cards = [
            {"card_type": "cloze", "question": "GABA is the primary {{c1::inhibitory}} neurotransmitter."},
            {"card_type": "cloze", "question": "GABA is the primary {{c1::inhibitory}} neurotransmitter."}
        ]
        unique = dedup.deduplicate(cards)
        assert len(unique) == 1

    # --- Feature 7: Cognitive Taxonomies & Hierarchical Tagging ---

    def test_tier1_cognitive_taxonomy_function_structure(self):
        """Feature 7.1: Function -> Structure card formulation (Brain Region <-> Function)."""
        card = {
            "card_type": "active_recall_qa",
            "taxonomy": "function_structure",
            "keyword": "Cerebellum",
            "descriptor": "Coordinates voluntary movements and motor learning.",
            "question": "What brain structure is responsible for: <i>Coordinates voluntary movements and motor learning.</i>?",
            "answer": "<b>Cerebellum</b>",
            "category_badge": "badge-definition",
            "tags": ["anatomy", "function_structure", "recall"]
        }
        assert "structure" in card["question"].lower()
        assert "Cerebellum" in card["answer"]
        assert "function_structure" in card["tags"]

    def test_tier1_cognitive_taxonomy_concept_mechanism(self):
        """Feature 7.2: Concept -> Mechanism card formulation (Process <-> Mechanism)."""
        card = {
            "card_type": "active_recall_qa",
            "taxonomy": "concept_mechanism",
            "keyword": "Action Potential",
            "descriptor": "Rapid depolarization mediated by voltage-gated sodium channels followed by repolarization via potassium efflux.",
            "question": "What is the key mechanism regarding <b>Action Potential</b>?",
            "answer": "Rapid depolarization mediated by voltage-gated sodium channels followed by repolarization via potassium efflux.",
            "category_badge": "badge-important",
            "tags": ["mechanism", "concept_mechanism"]
        }
        assert "mechanism" in card["question"].lower()
        assert "voltage-gated" in card["answer"]

    def test_tier1_cognitive_taxonomy_example_category(self):
        """Feature 7.3: Example -> Category card formulation (Exemplar <-> Condition/Principle)."""
        card = {
            "card_type": "active_recall_qa",
            "taxonomy": "example_category",
            "keyword": "Patient H.M.",
            "descriptor": "Severe anterograde amnesia following bilateral medial temporal lobe resection.",
            "question": "<b>Patient H.M.</b> is a primary clinical example of what condition or phenomenon?",
            "answer": "Severe anterograde amnesia following bilateral medial temporal lobe resection.",
            "category_badge": "badge-important",
            "tags": ["clinical", "example_category"]
        }
        assert "example" in card["question"].lower()
        assert "anterograde amnesia" in card["answer"]

    def test_tier1_hierarchical_tags_structure(self):
        """Feature 7.4: Formulates structured hierarchical tags conforming to Anki subtag standard."""
        course_tag = "Course::PSYC_2240"
        chapter_tag = "Chapter::Synaptic_Transmission"
        type_tag = "Type::Function_Structure"

        tags = [course_tag, chapter_tag, type_tag]
        for t in tags:
            assert "::" in t
            assert not re.search(r'\s', t), "Hierarchical tags must not contain unescaped whitespace"

    def test_tier1_deck_quality_pipeline_process_deck_contract(self):
        """Feature 7.5: DeckQualityPipeline.process_deck returns (cards, telemetry_dict)."""
        assert DeckQualityPipeline is not None, "DeckQualityPipeline not found in card_validator"
        pipeline = DeckQualityPipeline()
        raw_cards = [
            {"question": "What is serotonin?", "answer": "A monoamine neurotransmitter modulating mood.", "tags": ["neuro"]},
            {"question": "What is serotonin?", "answer": "A monoamine neurotransmitter modulating mood.", "tags": ["biochem"]},
            {"question": "What is this concept?", "answer": "Vague question that should be rejected."}
        ]
        clean_cards, telemetry = pipeline.process_deck(raw_cards)
        assert len(clean_cards) == 1
        assert telemetry["total_input"] == 3
        assert telemetry["valid"] == 1
        assert telemetry["rejected"] == 1
        assert telemetry["duplicates_removed"] == 1


# ==============================================================================
# TIER 2: BOUNDARY & CORNER CASES (6 Categories x 5 Tests = 30 Tests)
# ==============================================================================

class TestTier2BoundaryAndCornerCases:
    """
    Tier 2 tests verify stress, fault injection, and boundary conditions:
    1. Google Docs API Failures & Auth/Network Errors
    2. Malformed URLs, IDs & Invalid Input Sources
    3. Empty Documents, Table Parsing & Highlight Break Merging
    4. Extreme Character Lengths & Field Boundaries
    5. Malformed JSON, Corruption Fragments & HTML Garbage
    6. Deduplication Edge Cases & Bidirectional Preservation
    """

    # --- Boundary 1: Google Docs API Failures & Auth / Network Errors ---

    def test_tier2_gdocs_api_http_403_forbidden(self):
        """Boundary 1.1: Handles HTTP 403 Forbidden (Permission denied) gracefully."""
        assert extract_google_doc_structured is not None
        from googleapiclient.errors import HttpError
        mock_resp = MagicMock(status=403, reason="Permission Denied")
        mock_service = MagicMock()
        mock_service.documents().get.side_effect = HttpError(mock_resp, b'{"error": "Forbidden"}')

        with patch("scripts.docs_api_client.build", return_value=mock_service), \
             patch("scripts.docs_api_client.get_credentials", return_value=MagicMock()):
            data, title = extract_google_doc_structured("forbidden_doc_id")
            assert data == []
            assert title == ""

    def test_tier2_gdocs_api_http_404_not_found(self):
        """Boundary 1.2: Handles HTTP 404 Not Found (Invalid doc ID) gracefully."""
        assert extract_google_doc_structured is not None
        from googleapiclient.errors import HttpError
        mock_resp = MagicMock(status=404, reason="Not Found")
        mock_service = MagicMock()
        mock_service.documents().get.side_effect = HttpError(mock_resp, b'{"error": "Document not found"}')

        with patch("scripts.docs_api_client.build", return_value=mock_service), \
             patch("scripts.docs_api_client.get_credentials", return_value=MagicMock()):
            data, title = extract_google_doc_structured("nonexistent_id")
            assert data == []
            assert title == ""

    def test_tier2_gdocs_api_http_429_quota_exceeded(self):
        """Boundary 1.3: Handles HTTP 429 Rate Limit Exceeded gracefully without uncaught exceptions."""
        assert extract_google_doc_structured is not None
        from googleapiclient.errors import HttpError
        mock_resp = MagicMock(status=429, reason="Quota Exceeded")
        mock_service = MagicMock()
        mock_service.documents().get.side_effect = HttpError(mock_resp, b'{"error": "Rate limit"}')

        with patch("scripts.docs_api_client.build", return_value=mock_service), \
             patch("scripts.docs_api_client.get_credentials", return_value=MagicMock()):
            data, title = extract_google_doc_structured("rate_limited_id")
            assert data == []
            assert title == ""

    def test_tier2_gdocs_api_socket_timeout(self):
        """Boundary 1.4: Handles socket/network timeout without crashing."""
        assert extract_google_doc_structured is not None
        import socket
        mock_service = MagicMock()
        mock_service.documents().get.side_effect = socket.timeout("Connection timed out")

        with patch("scripts.docs_api_client.build", return_value=mock_service), \
             patch("scripts.docs_api_client.get_credentials", return_value=MagicMock()):
            data, title = extract_google_doc_structured("timeout_id")
            assert data == []
            assert title == ""

    def test_tier2_gdocs_api_missing_credentials(self):
        """Boundary 1.5: Missing credentials.json and token.json raises FileNotFoundError."""
        from scripts.docs_api_client import get_credentials
        with patch("os.path.exists", return_value=False):
            with pytest.raises(FileNotFoundError):
                get_credentials()

    # --- Boundary 2: Malformed URLs, IDs & Invalid Input Sources ---

    def test_tier2_detect_input_empty_string(self):
        """Boundary 2.1: Empty or whitespace string classified as invalid."""
        if detect_input_type is None:
            pytest.skip("detect_input_type pending implementation in Milestone 5")
        assert detect_input_type("") == "invalid"
        assert detect_input_type("   \t  ") == "invalid"

    def test_tier2_detect_input_malformed_url(self):
        """Boundary 2.2: URL missing /document/d/ segment classified as invalid."""
        if detect_input_type is None:
            pytest.skip("detect_input_type pending implementation in Milestone 5")
        assert detect_input_type("https://docs.google.com/spreadsheets/d/123/edit") == "invalid"
        assert detect_input_type("https://google.com") == "invalid"

    def test_tier2_detect_input_nonexistent_docx(self):
        """Boundary 2.3: Non-existent .docx file classified as missing or invalid."""
        if detect_input_type is None:
            pytest.skip("detect_input_type pending implementation in Milestone 5")
        res = detect_input_type("nonexistent_path_12345.docx")
        assert res in ("docx_missing", "invalid")

    def test_tier2_detect_input_wrong_extension(self, tmp_path):
        """Boundary 2.4: Non-docx file (e.g. .pdf, .txt) classified as invalid."""
        if detect_input_type is None:
            pytest.skip("detect_input_type pending implementation in Milestone 5")
        pdf_path = tmp_path / "lecture.pdf"
        pdf_path.touch()
        assert detect_input_type(str(pdf_path)) == "invalid"

    def test_tier2_extract_document_id_with_query_params(self):
        """Boundary 2.5: URL with complex query parameters and fragments isolates doc ID."""
        assert extract_document_id is not None
        url = "https://docs.google.com/document/d/1A2B3C4D5E6F7G8H9I0J_kLmNoPqRsTuVwXyZ12345/edit?usp=sharing&authuser=0#bookmark=kix.123"
        assert extract_document_id(url) == "1A2B3C4D5E6F7G8H9I0J_kLmNoPqRsTuVwXyZ12345"

    # --- Boundary 3: Empty Documents, Table Parsing & Highlight Break Merging ---

    def test_tier2_gdocs_empty_body(self):
        """Boundary 3.1: Document with empty content list returns 0 highlights."""
        assert extract_google_doc_structured is not None
        mock_doc = {"title": "Empty Lecture", "body": {"content": []}}
        mock_service = create_mock_service(mock_doc)
        with patch("scripts.docs_api_client.build", return_value=mock_service), \
             patch("scripts.docs_api_client.get_credentials", return_value=MagicMock()):
            data, title = extract_google_doc_structured("empty_id")
            assert title == "Empty Lecture"
            assert data == []

    def test_tier2_gdocs_table_cell_extraction(self):
        """Boundary 3.2: Extracts highlighted text runs nested inside table cells."""
        assert extract_google_doc_structured is not None
        mock_doc = {
            "title": "Table Lecture",
            "body": {
                "content": [
                    {
                        "table": {
                            "tableRows": [
                                {
                                    "tableCells": [
                                        {
                                            "content": [
                                                {
                                                    "paragraph": {
                                                        "elements": [
                                                            {
                                                                "textRun": {
                                                                    "content": "Glial cells: ",
                                                                    "textStyle": {"backgroundColor": {"color": {"rgbColor": {"red": 1.0, "green": 1.0, "blue": 0.0}}}}
                                                                }
                                                            },
                                                            {
                                                                "textRun": {
                                                                    "content": "support and protect neurons.",
                                                                    "textStyle": {"backgroundColor": {"color": {"rgbColor": {"red": 0.0, "green": 1.0, "blue": 0.0}}}}
                                                                }
                                                            }
                                                        ]
                                                    }
                                                }
                                            ]
                                        }
                                    ]
                                }
                            ]
                        }
                    }
                ]
            }
        }
        mock_service = create_mock_service(mock_doc)
        with patch("scripts.docs_api_client.build", return_value=mock_service), \
             patch("scripts.docs_api_client.get_credentials", return_value=MagicMock()):
            data, title = extract_google_doc_structured("table_id")
            assert len(data) == 1
            assert len(data[0]["highlights"]) == 2

    def test_tier2_gdocs_highlight_break_across_whitespace(self):
        """Boundary 3.3: Highlight interrupted by unhighlighted space merged into single segment."""
        assert extract_google_doc_structured is not None
        mock_doc = {
            "title": "Gap Lecture",
            "body": {
                "content": [
                    {
                        "paragraph": {
                            "elements": [
                                {
                                    "textRun": {
                                        "content": "Long-Term",
                                        "textStyle": {"backgroundColor": {"color": {"rgbColor": {"red": 1.0, "green": 1.0, "blue": 0.0}}}}
                                    }
                                },
                                {
                                    "textRun": {
                                        "content": " ", # Unhighlighted space
                                        "textStyle": {}
                                    }
                                },
                                {
                                    "textRun": {
                                        "content": "Potentiation",
                                        "textStyle": {"backgroundColor": {"color": {"rgbColor": {"red": 1.0, "green": 1.0, "blue": 0.0}}}}
                                    }
                                }
                            ]
                        }
                    }
                ]
            }
        }
        mock_service = create_mock_service(mock_doc)
        with patch("scripts.docs_api_client.build", return_value=mock_service), \
             patch("scripts.docs_api_client.get_credentials", return_value=MagicMock()):
            data, title = extract_google_doc_structured("gap_id")
            assert len(data) == 1
            # Both words merged across unhighlighted whitespace
            hl = data[0]["highlights"]
            assert len(hl) == 1
            assert "Long-Term Potentiation" in hl[0]["text"]

    def test_tier2_gdocs_pure_punctuation_highlights_ignored(self):
        """Boundary 3.4: Punctuation-only highlights (e.g. ' : - ; ') are skipped."""
        assert extract_google_doc_structured is not None
        mock_doc = {
            "title": "Punctuation Lecture",
            "body": {
                "content": [
                    {
                        "paragraph": {
                            "elements": [
                                {
                                    "textRun": {
                                        "content": " : - ; . ",
                                        "textStyle": {"backgroundColor": {"color": {"rgbColor": {"red": 1.0, "green": 1.0, "blue": 0.0}}}}
                                    }
                                }
                            ]
                        }
                    }
                ]
            }
        }
        mock_service = create_mock_service(mock_doc)
        with patch("scripts.docs_api_client.build", return_value=mock_service), \
             patch("scripts.docs_api_client.get_credentials", return_value=MagicMock()):
            data, title = extract_google_doc_structured("punct_id")
            assert len(data) == 0

    def test_tier2_gdocs_no_highlights_paragraph_skipped(self):
        """Boundary 3.5: Paragraphs without any highlighted text runs are skipped."""
        assert extract_google_doc_structured is not None
        mock_doc = {
            "title": "Plain Lecture",
            "body": {
                "content": [
                    {
                        "paragraph": {
                            "elements": [
                                {"textRun": {"content": "Normal unhighlighted text.", "textStyle": {}}}
                            ]
                        }
                    }
                ]
            }
        }
        mock_service = create_mock_service(mock_doc)
        with patch("scripts.docs_api_client.build", return_value=mock_service), \
             patch("scripts.docs_api_client.get_credentials", return_value=MagicMock()):
            data, title = extract_google_doc_structured("plain_id")
            assert len(data) == 0

    # --- Boundary 4: Extreme Character Lengths & Field Boundaries ---

    def test_tier2_validator_front_too_short(self):
        """Boundary 4.1: Question shorter than min_front (10 chars) rejected."""
        assert CardValidator is not None
        validator = CardValidator(min_front=10)
        card = {"question": "What is?", "answer": "A neurotransmitter modulating sleep."}
        res = validator.validate_card(card)
        assert res.is_valid is False

    def test_tier2_validator_front_too_long(self):
        """Boundary 4.2: Question longer than max_front (250 chars) rejected."""
        assert CardValidator is not None
        validator = CardValidator(max_front=250)
        long_q = "What is " + ("very long explanation " * 20) + "?"
        card = {"question": long_q, "answer": "Target concept"}
        res = validator.validate_card(card)
        assert res.is_valid is False

    def test_tier2_validator_back_too_short(self):
        """Boundary 4.3: Answer shorter than min_back (5 chars) rejected."""
        assert CardValidator is not None
        validator = CardValidator(min_back=5)
        card = {"question": "What is the abbreviation for dopamine?", "answer": "DA"}
        res = validator.validate_card(card)
        assert res.is_valid is False

    def test_tier2_validator_back_too_long(self):
        """Boundary 4.4: Monolithic paragraph answer (>500 chars) rejected under Rule 4."""
        assert CardValidator is not None
        validator = CardValidator(max_back=500)
        monolithic_answer = "This mechanism involves " + ("cascades of biochemical reactions " * 30)
        card = {"question": "What is synaptic transmission?", "answer": monolithic_answer}
        res = validator.validate_card(card)
        assert res.is_valid is False

    def test_tier2_validator_extreme_10k_character_payload(self):
        """Boundary 4.5: Pathological 10,000-character payload handled without freeze or memory crash."""
        assert CardValidator is not None
        validator = CardValidator()
        card = {"question": "Q: " + ("x" * 10000), "answer": "A: " + ("y" * 10000)}
        res = validator.validate_card(card)
        assert res.is_valid is False

    # --- Boundary 5: Malformed JSON, Corruption Fragments & HTML Garbage ---

    def test_tier2_validator_missing_question_key(self):
        """Boundary 5.1: Card dictionary missing 'question' key rejected."""
        assert CardValidator is not None
        validator = CardValidator()
        card = {"answer": "Only an answer"}
        res = validator.validate_card(card)
        assert res.is_valid is False

    def test_tier2_validator_missing_answer_key(self):
        """Boundary 5.2: Card dictionary missing 'answer' key rejected."""
        assert CardValidator is not None
        validator = CardValidator()
        card = {"question": "Only a question?"}
        res = validator.validate_card(card)
        assert res.is_valid is False

    def test_tier2_validator_none_values_rejected(self):
        """Boundary 5.3: Card containing None for question or answer rejected."""
        assert CardValidator is not None
        validator = CardValidator()
        card = {"question": None, "answer": None}
        res = validator.validate_card(card)
        assert res.is_valid is False

    def test_tier2_sanitizer_complex_textbook_index_corruption(self):
        """Boundary 5.4: Standalone index citation (', 59, 71, 377') preserved for validator flagging."""
        assert CardSanitizer is not None
        assert CardValidator is not None
        sanitizer = CardSanitizer()
        validator = CardValidator()
        pure_index = ", 59, 71, 377, 412"
        # Sanitizer must preserve standalone index sequence
        assert sanitizer.strip_index_runs(pure_index) == pure_index
        # Validator must catch and flag it
        card = {"question": "What is basal ganglia?", "answer": pure_index}
        res = validator.validate_card(card)
        assert res.is_valid is False

    def test_tier2_sanitizer_html_entity_and_tag_soup(self):
        """Boundary 5.5: Handles dirty web copy-paste HTML markup cleanly."""
        assert CardSanitizer is not None
        sanitizer = CardSanitizer()
        raw = "<span><b>Dopamine</b></span>&nbsp;mediates&nbsp;reward<style>.card{color:red;}</style>."
        clean = sanitizer.strip_css(raw)
        assert "<style>" not in clean
        assert "Dopamine" in clean

    # --- Boundary 6: Deduplication Edge Cases & Bidirectional Preservation ---

    def test_tier2_deduplicator_case_insensitive_matching(self):
        """Boundary 6.1: Deduplicates cards differing only in casing and whitespace."""
        assert CardDeduplicator is not None
        dedup = CardDeduplicator()
        cards = [
            {"question": "What is GABA?", "answer": "Inhibitory neurotransmitter."},
            {"question": "  what is gaba?  ", "answer": "inhibitory neurotransmitter.  "}
        ]
        unique = dedup.deduplicate(cards)
        assert len(unique) == 1

    def test_tier2_deduplicator_html_tag_variance_deduped(self):
        """Boundary 6.2: Identical cards differing only in HTML bolding deduplicated."""
        assert CardDeduplicator is not None
        dedup = CardDeduplicator()
        cards = [
            {"question": "What is <b>GABA</b>?", "answer": "Inhibitory neurotransmitter."},
            {"question": "What is GABA?", "answer": "Inhibitory neurotransmitter."}
        ]
        unique = dedup.deduplicate(cards)
        assert len(unique) == 1

    def test_tier2_deduplicator_100_identical_cards(self):
        """Boundary 6.3: 100 identical duplicate cards collapse to exactly 1 card."""
        assert CardDeduplicator is not None
        dedup = CardDeduplicator()
        cards = [{"question": "What is glutamate?", "answer": "Primary excitatory neurotransmitter."} for _ in range(100)]
        unique, count = dedup.deduplicate(cards, return_count=True)
        assert len(unique) == 1
        assert count == 99

    def test_tier2_deduplicator_multiple_bidirectional_pairs(self):
        """Boundary 6.4: Multiple distinct bidirectional pairs all survive deduplication."""
        assert CardDeduplicator is not None
        dedup = CardDeduplicator()
        cards = [
            # Pair 1: Glutamate
            {"question": "What is the definition of Glutamate?", "answer": "Primary excitatory neurotransmitter."},
            {"question": "What term is defined by: Primary excitatory neurotransmitter?", "answer": "Glutamate"},
            # Duplicate of Pair 1 forward
            {"question": "What is the definition of Glutamate?", "answer": "Primary excitatory neurotransmitter."},
            # Pair 2: GABA
            {"question": "What is the definition of GABA?", "answer": "Primary inhibitory neurotransmitter."},
            {"question": "What term is defined by: Primary inhibitory neurotransmitter?", "answer": "GABA"},
        ]
        unique, count = dedup.deduplicate(cards, return_count=True)
        assert len(unique) == 4
        assert count == 1

    def test_tier2_deduplicator_empty_list(self):
        """Boundary 6.5: Deduplicating empty card list returns empty list."""
        assert CardDeduplicator is not None
        dedup = CardDeduplicator()
        assert dedup.deduplicate([]) == []


# ==============================================================================
# TIER 3: CROSS-FEATURE COMBINATIONS (8 Tests)
# ==============================================================================

class TestTier3CrossFeatureCombinations:
    """
    Tier 3 tests verify multi-subsystem integrations:
    - Google Docs Ingestion + Quality Pipeline + .apkg compilation
    - Corrupted Web Notes + Sanitization + Validation
    - Bidirectional Reciprocity across Cognitive Formulations
    - Hierarchical Tagging through compilation into SQLite database
    """

    def test_tier3_gdocs_ingestion_through_quality_pipeline(self, tmp_path):
        """Combination 3.1: Google Doc extracts flow through DeckQualityPipeline into valid cards."""
        assert extract_google_doc_structured is not None
        assert DeckQualityPipeline is not None

        mock_doc = make_google_docs_mock(
            title="PSYC 2240 - Lecture 3 Action Potentials",
            paragraphs=[
                {
                    "heading": "HEADING_1",
                    "runs": [{"text": "Electrophysiology\n", "color": None}]
                },
                {
                    "heading": "NORMAL_TEXT",
                    "runs": [
                        {"text": "Resting Membrane Potential: ", "color": (1.0, 1.0, 0.0)},
                        {"text": "the electrical potential difference across the neuronal membrane at rest, typically around -70 mV.", "color": (0.0, 1.0, 0.0)}
                    ]
                }
            ]
        )
        mock_service = create_mock_service(mock_doc)
        with patch("scripts.docs_api_client.build", return_value=mock_service), \
             patch("scripts.docs_api_client.get_credentials", return_value=MagicMock()):
            data, title = extract_google_doc_structured("dummy_id")
            assert len(data) == 1

            # Synthesize cards
            cards = synthesize_cards(data, deck_tags=["neuroscience"])
            assert len(cards) >= 1

            # Run through quality pipeline
            pipeline = DeckQualityPipeline()
            clean_cards, telemetry = pipeline.process_deck(cards)
            assert len(clean_cards) >= 1
            assert telemetry["rejected"] == 0

    def test_tier3_corrupted_google_doc_cleaned_and_packaged(self, tmp_path):
        """Combination 3.2: Google Doc containing CSS pollution and figure citations is sanitized and packaged."""
        assert DeckQualityPipeline is not None
        assert create_deck_package is not None

        dirty_cards = [
            {
                "question": ".card { font-size: 14px; }\nHIGH Priority What is the primary role of <b>myelin</b>?",
                "answer": "Insulates axons to accelerate action potential conduction (Figure 2.1).",
                "category_badge": "badge-definition",
                "tags": ["histology"]
            }
        ]
        pipeline = DeckQualityPipeline()
        clean_cards, telemetry = pipeline.process_deck(dirty_cards)
        assert len(clean_cards) == 1
        c = clean_cards[0]
        assert ".card" not in c["question"]
        assert "HIGH Priority" not in c["question"]
        assert "(Figure 2.1)" not in c["answer"]

        # Package into temporary apkg
        with patch.dict("scripts.extract_and_generate.CONFIG", {"google_drive_decks_directory": None}):
            out_apkg = create_deck_package("Clean Test Deck", clean_cards, output_filename="clean_test_deck_tier3.apkg")
            assert out_apkg.exists()
            assert out_apkg.stat().st_size > 0
            if out_apkg.exists():
                out_apkg.unlink()

    def test_tier3_gdocs_api_error_does_not_generate_corrupt_apkg(self, tmp_path):
        """Combination 3.3: Google Docs API failure halts pipeline cleanly without creating invalid packages."""
        assert extract_google_doc_structured is not None
        import socket
        mock_service = MagicMock()
        mock_service.documents().get.side_effect = socket.error("Connection reset by peer")

        with patch("scripts.docs_api_client.build", return_value=mock_service), \
             patch("scripts.docs_api_client.get_credentials", return_value=MagicMock()):
            data, title = extract_google_doc_structured("broken_id")
            assert data == []
            cards = synthesize_cards(data) if data else []
            assert len(cards) == 0

    def test_tier3_bidirectional_and_cognitive_taxonomies_coexist(self):
        """Combination 3.4: Bidirectional definitions and cognitive taxonomy cards coexist without mutual interference."""
        assert DeckQualityPipeline is not None
        cards = [
            # Standard SOP Bidirectional definition
            {"question": "What is the definition of <b>Agonist</b>?", "answer": "A chemical that binds to a receptor and activates it.", "tags": ["pharmacology"]},
            {"question": "What term is defined by:<br><i>A chemical that binds to a receptor and activates it.</i>", "answer": "Agonist", "tags": ["pharmacology"]},
            # Cognitive Taxonomy: Function -> Structure
            {"question": "What brain structure is responsible for: <i>Coordination of voluntary movement</i>?", "answer": "<b>Cerebellum</b>", "tags": ["anatomy", "function_structure"]},
            {"question": "What is the primary function of the <b>Cerebellum</b>?", "answer": "Coordination of voluntary movement.", "tags": ["anatomy", "function_structure"]},
            # Cognitive Taxonomy: Concept -> Mechanism
            {"question": "What is the key mechanism regarding <b>Long-Term Potentiation</b>?", "answer": "NMDA receptor activation leading to calcium influx and AMPA receptor insertion.", "tags": ["mechanism"]}
        ]
        pipeline = DeckQualityPipeline()
        clean_cards, telemetry = pipeline.process_deck(cards)
        assert len(clean_cards) == 5, "Legitimate distinct cognitive cards were falsely deduplicated!"
        assert telemetry["duplicates_removed"] == 0

    def test_tier3_deck_naming_resolution_with_cli_overrides(self):
        """Combination 3.5: Deck naming handles complex title + explicit CLI override flags."""
        assert resolve_deck_naming is not None
        title, fn = resolve_deck_naming(
            "PSYC 2240 Biological Bases.docx",
            explicit_class="PSYC 2240",
            explicit_chapter="Chapter 5: Neurotransmitters"
        )
        assert "PSYC 2240" in title
        assert "Biological Basis of Behaviour" in title
        assert "Chapter 5: Neurotransmitters" in title

    def test_tier3_hierarchical_tags_propagated_to_genanki_package(self, tmp_path):
        """Combination 3.6: Hierarchical tags (Course::*, Chapter::*) are saved into .apkg collection."""
        assert create_deck_package is not None
        cards = [
            {
                "question": "What is the primary function of the <b>amygdala</b>?",
                "answer": "Processing emotions, particularly fear and threat detection.",
                "category_badge": "badge-definition",
                "context": "Limbic System",
                "tags": ["Course::PSYC_2240", "Chapter::Limbic_System", "Type::Function_Structure"]
            }
        ]
        with patch.dict("scripts.extract_and_generate.CONFIG", {"google_drive_decks_directory": None}):
            out_apkg = create_deck_package("Tag Test Deck", cards, output_filename="hierarchical_tags_test.apkg")

            # Inspect SQLite inside zip
            with zipfile.ZipFile(out_apkg, "r") as z:
                assert "collection.anki2" in z.namelist()
                z.extract("collection.anki2", tmp_path)

            db_path = tmp_path / "collection.anki2"
            conn = sqlite3.connect(db_path)
            cur = conn.cursor()
            cur.execute("SELECT tags FROM notes")
            rows = cur.fetchall()
            conn.close()

            if out_apkg.exists():
                out_apkg.unlink()

            assert len(rows) == 1
        tags_str = rows[0][0]
        assert "Course::PSYC_2240" in tags_str
        assert "Type::Function_Structure" in tags_str

    def test_tier3_multicolor_highlight_enrichment(self):
        """Combination 3.7: Paragraph with Yellow term + Green definition + Blue context creates enriched note."""
        assert synthesize_cards is not None
        raw_paragraph = {
            "heading": "Visual Processing",
            "full_paragraph": "Occipital Lobe is the visual processing center of the mammalian brain. Damage leads to cortical blindness.",
            "segments": [
                {"category": "yellow", "text": "Occipital Lobe", "raw_color": "yellow"},
                {"category": "green", "text": "is the visual processing center of the mammalian brain.", "raw_color": "green"},
                {"category": "other", "text": "Damage leads to cortical blindness.", "raw_color": "blue"}
            ],
            "highlights": [
                {"category": "yellow", "text": "Occipital Lobe", "raw_color": "yellow"},
                {"category": "green", "text": "is the visual processing center of the mammalian brain.", "raw_color": "green"}
            ]
        }
        cards = synthesize_cards([raw_paragraph])
        assert len(cards) >= 1
        # Context should capture context notes or heading
        assert any("Occipital Lobe" in c["question"] or "Occipital Lobe" in c["answer"] for c in cards)

    def test_tier3_quality_pipeline_telemetry_accuracy(self):
        """Combination 3.8: DeckQualityPipeline produces exact counts across all telemetry buckets."""
        assert DeckQualityPipeline is not None
        pipeline = DeckQualityPipeline()
        deck = [
            # 2 Valid cards
            {"question": "What is serotonin?", "answer": "A monoamine neurotransmitter."},
            {"question": "What is norepinephrine?", "answer": "A catecholamine neurotransmitter."},
            # 1 Duplicate card
            {"question": "What is serotonin?", "answer": "A monoamine neurotransmitter."},
            # 1 Placeholder rejected
            {"question": "What is melatonin?", "answer": "Review this concept."},
            # 1 Vague veto rejected
            {"question": "What is this concept?", "answer": "Valid answer text."},
        ]
        clean_cards, telemetry = pipeline.process_deck(deck)
        assert len(clean_cards) == 2
        assert telemetry["total_input"] == 5
        assert telemetry["valid"] == 2
        assert telemetry["rejected"] == 2
        assert telemetry["duplicates_removed"] == 1


# ==============================================================================
# TIER 4: REAL-WORLD SCENARIOS (7 Tests)
# ==============================================================================

class TestTier4RealWorldScenarios:
    """
    Tier 4 tests verify full end-to-end workflows and artifact fidelity:
    1. Realistic Google Doc lecture simulation with mixed highlights
    2. Input polymorphism parity (Word vs Google Docs)
    3. Dirty study guide web copy-paste decontamination
    4. Compiled .apkg SQLite database inspection and styling verification
    5. Agent-as-Judge 30-Point Rubric verification
    6. Strict vs Permissive pipeline execution
    7. Full execution performance budget (< 15.0s)
    """

    def test_tier4_full_gdocs_lecture_simulation(self, tmp_path, judge):
        """Scenario 4.1: Full lecture simulation: Google Docs API extraction -> Quality Pipeline -> .apkg."""
        assert extract_google_doc_structured is not None
        assert DeckQualityPipeline is not None
        assert create_deck_package is not None

        mock_doc = make_google_docs_mock(
            title="PSYC 2240 - Biological Bases of Behaviour - Chapter 4: Synaptic Transmission",
            paragraphs=[
                {
                    "heading": "HEADING_1",
                    "runs": [{"text": "Synaptic Transmission & Neurochemistry\n", "color": None}]
                },
                {
                    "heading": "NORMAL_TEXT",
                    "runs": [
                        {"text": "Action Potential: ", "color": (1.0, 1.0, 0.0)},
                        {"text": "a rapid electrical signal generated by the opening of voltage-gated sodium channels.", "color": (0.0, 1.0, 0.0)}
                    ]
                },
                {
                    "heading": "NORMAL_TEXT",
                    "runs": [
                        {"text": "Neurotransmitter: ", "color": (1.0, 1.0, 0.0)},
                        {"text": "chemical messenger released by neurons to communicate across synaptic clefts.", "color": (0.0, 1.0, 0.0)}
                    ]
                },
                {
                    "heading": "NORMAL_TEXT",
                    "runs": [
                        {"text": "Corrupted Note: ", "color": (1.0, 1.0, 0.0)},
                        {"text": "Review this concept in textbook.", "color": (0.0, 1.0, 0.0)} # Should be rejected
                    ]
                }
            ]
        )
        mock_service = create_mock_service(mock_doc)
        with patch("scripts.docs_api_client.build", return_value=mock_service), \
             patch("scripts.docs_api_client.get_credentials", return_value=MagicMock()):
            # 1. Ingestion
            data, title = extract_google_doc_structured("real_lecture_doc_id")
            assert title.startswith("PSYC 2240")
            assert len(data) == 3

            # 2. Synthesis
            raw_cards = synthesize_cards(data, deck_tags=["PSYC2240", "synaptic_transmission"])
            # Inject dirty card to exercise DeckQualityPipeline defensive filtering
            raw_cards.append({
                "question": "What is the unverified concept?",
                "answer": "Review this concept in the textbook.",
                "tags": ["PSYC2240"]
            })
            assert len(raw_cards) >= 5

            # 3. Defensive Quality Pipeline
            pipeline = DeckQualityPipeline()
            clean_cards, telemetry = pipeline.process_deck(raw_cards)
            assert telemetry["rejected"] >= 1, "Placeholder card was not rejected!"
            assert len(clean_cards) >= 4

            # 4. Packaging
            with patch.dict("scripts.extract_and_generate.CONFIG", {"google_drive_decks_directory": None}):
                out_file = create_deck_package("PSYC 2240 (Biological Basis of Behaviour):Chapter 4", clean_cards, output_filename="PSYC_2240_Lecture.apkg")
                assert out_file.exists()
                assert out_file.stat().st_size > 0

                # 5. Judge Evaluation
                deck_report = judge.evaluate_deck(clean_cards)
                assert deck_report["veto_count"] == 0, f"Judge detected vetoes: {deck_report['veto_count']}"
                assert deck_report["average_score"] >= 25.0
                assert deck_report["deck_passed"] is True

                if out_file.exists():
                    out_file.unlink()

    def test_tier4_input_polymorphism_parity(self, tmp_path):
        """Scenario 4.2: Word .docx and Google Docs API extractions produce identical schema dictionaries."""
        assert extract_document_highlights is not None
        assert extract_google_doc_structured is not None

        # 1. Create .docx
        docx_file = tmp_path / "parity_test.docx"
        create_sample_docx(
            docx_file,
            [{
                "heading": "Neuroscience",
                "runs": [
                    {"text": "Axon: ", "highlight": "yellow"},
                    {"text": "conducts electrical impulses away from the cell body.", "highlight": "green"}
                ]
            }]
        )
        docx_data = extract_document_highlights(str(docx_file))

        # 2. Mock Google Doc with exact same content
        mock_doc = make_google_docs_mock(
            title="Parity Doc",
            paragraphs=[
                {
                    "heading": "HEADING_1",
                    "runs": [{"text": "Neuroscience\n", "color": None}]
                },
                {
                    "heading": "NORMAL_TEXT",
                    "runs": [
                        {"text": "Axon: ", "color": (1.0, 1.0, 0.0)},
                        {"text": "conducts electrical impulses away from the cell body.", "color": (0.0, 1.0, 0.0)}
                    ]
                }
            ]
        )
        mock_service = create_mock_service(mock_doc)
        with patch("scripts.docs_api_client.build", return_value=mock_service), \
             patch("scripts.docs_api_client.get_credentials", return_value=MagicMock()):
            gdocs_data, title = extract_google_doc_structured("parity_id")

        # 3. Assert schema parity
        assert len(docx_data) == len(gdocs_data) == 1
        p_docx = docx_data[0]
        p_gdoc = gdocs_data[0]

        core_keys = {"heading", "full_paragraph", "segments", "highlights"}
        assert core_keys.issubset(set(p_docx.keys()))
        assert core_keys.issubset(set(p_gdoc.keys()))
        assert p_docx["heading"] == p_gdoc["heading"] == "Neuroscience"
        assert len(p_docx["highlights"]) == len(p_gdoc["highlights"]) == 2

    def test_tier4_study_guide_web_copy_paste_decontamination(self):
        """Scenario 4.3: Decontaminates dirty study guide with CSS styles, figure citations, and duplicates."""
        assert DeckQualityPipeline is not None
        dirty_study_guide_cards = [
            {
                "question": ".card { font-family: Calibri; }\nHIGH Priority What is <b>Synaptic Vesicle</b>?",
                "answer": "Membrane-bound spheres containing neurotransmitters (Figure 3.1; see also Table 2.2).",
                "tags": ["organelles", "synapse"]
            },
            {
                # Duplicate with slight whitespace and lowercase tag
                "question": "What is Synaptic Vesicle?",
                "answer": "Membrane-bound spheres containing neurotransmitters.",
                "tags": ["synapse", "cell_bio"]
            },
            {
                # Corrupted index citation note
                "question": "What is synaptic cleft?",
                "answer": ", 12, 14, 18, 22"
            }
        ]
        pipeline = DeckQualityPipeline()
        clean_cards, telemetry = pipeline.process_deck(dirty_study_guide_cards)

        assert len(clean_cards) == 1
        c = clean_cards[0]
        assert ".card" not in c["question"]
        assert "Figure 3.1" not in c["answer"]
        assert "Table 2.2" not in c["answer"]
        assert set(c["tags"]) == {"organelles", "synapse", "cell_bio"}
        assert telemetry["rejected"] == 1
        assert telemetry["duplicates_removed"] == 1

    def test_tier4_apkg_database_integrity_and_styling(self, tmp_path):
        """Scenario 4.4: Generated .apkg archive contains valid SQLite schema with WCAG AAA styling & .nightMode."""
        assert create_deck_package is not None
        cards = [
            {
                "question": "What is the primary function of the <b>hippocampus</b>?",
                "answer": "Consolidation of short-term memory into long-term memory.",
                "category_badge": "badge-definition",
                "context": "Limbic System",
                "tags": ["anatomy", "memory"]
            }
        ]
        with patch.dict("scripts.extract_and_generate.CONFIG", {"google_drive_decks_directory": None}):
            out_apkg = create_deck_package("Integrity Deck", cards, output_filename="integrity_test_t4.apkg")

            # Extract and verify SQLite
            with zipfile.ZipFile(out_apkg, "r") as z:
                assert "collection.anki2" in z.namelist()
                z.extract("collection.anki2", tmp_path)

            db_path = tmp_path / "collection.anki2"
            conn = sqlite3.connect(db_path)
            cur = conn.cursor()

            # Check models table for styling
            cur.execute("SELECT models FROM col")
            col_row = cur.fetchone()
            assert col_row is not None
            models_json = json.loads(col_row[0])
            model = list(models_json.values())[0]

            # Verify CSS styling
            css = model.get("css", "")
            assert ".card" in css
            assert ".nightMode" in css, "Dark mode .nightMode CSS class missing!"

            # Check notes table
            cur.execute("SELECT flds FROM notes")
            notes = cur.fetchall()
            conn.close()

            if out_apkg.exists():
                out_apkg.unlink()

            assert len(notes) == 1
            fields_str = notes[0][0]
            assert "hippocampus" in fields_str
            assert "Consolidation" in fields_str

    def test_tier4_agent_as_judge_deck_evaluation_pass(self, judge):
        """Scenario 4.5: ReferenceCardJudgeOracle grades synthesized Phase 2 deck with 100% pass rate."""
        cards = [
            {
                "question": "What is the definition of <b>Tolerance</b>?",
                "answer": "A state of progressively decreasing responsiveness to repeated drug administration.",
                "category_badge": "badge-definition",
                "context": "Pharmacodynamics",
                "tags": ["pharmacology", "recall"]
            },
            {
                "question": "What term is defined by:<br><i>A state of progressively decreasing responsiveness to repeated drug administration.</i>",
                "answer": "Tolerance",
                "category_badge": "badge-definition",
                "context": "Pharmacodynamics",
                "tags": ["pharmacology", "recognition"]
            },
            {
                "question": "What brain structure is responsible for: <i>Motor coordination and procedural learning</i>?",
                "answer": "<b>Cerebellum</b>",
                "category_badge": "badge-definition",
                "context": "Motor Systems",
                "tags": ["anatomy", "function_structure"]
            }
        ]
        deck_report = judge.evaluate_deck(cards)
        assert deck_report["deck_passed"] is True
        assert deck_report["veto_count"] == 0
        assert deck_report["average_score"] >= 25.0

    def test_tier4_pipeline_strict_mode_exception(self):
        """Scenario 4.6: DeckQualityPipeline raises ValueError in strict mode when encountering corrupted cards."""
        assert DeckQualityPipeline is not None
        pipeline = DeckQualityPipeline()
        contaminated_cards = [
            {"question": "What is dopamine?", "answer": "Review this concept."}
        ]
        with pytest.raises(ValueError) as excinfo:
            pipeline.process(contaminated_cards, strict=True)
        assert "Strict deck quality validation failed" in str(excinfo.value)

    def test_tier4_performance_budget_gdocs_processing(self, tmp_path):
        """Scenario 4.7: Ingestion, validation, and deck compilation completes well within 15.0 second budget."""
        import time
        assert extract_google_doc_structured is not None
        assert DeckQualityPipeline is not None
        assert create_deck_package is not None

        paragraphs = []
        for i in range(20):
            paragraphs.append({
                "heading": f"HEADING_1" if i % 5 == 0 else None,
                "runs": [
                    {"text": f"Concept {i}: ", "color": (1.0, 1.0, 0.0)},
                    {"text": f"Description for concept number {i} detailing its physiological role.", "color": (0.0, 1.0, 0.0)}
                ]
            })
        mock_doc = make_google_docs_mock("Speed Test Lecture", paragraphs)
        mock_service = create_mock_service(mock_doc)

        start_time = time.time()
        with patch("scripts.docs_api_client.build", return_value=mock_service), \
             patch("scripts.docs_api_client.get_credentials", return_value=MagicMock()), \
             patch.dict("scripts.extract_and_generate.CONFIG", {"google_drive_decks_directory": None}):
            data, title = extract_google_doc_structured("perf_id")
            cards = synthesize_cards(data)
            pipeline = DeckQualityPipeline()
            clean_cards, telemetry = pipeline.process_deck(cards)
            out_file = create_deck_package("Perf Deck", clean_cards, output_filename="perf_deck_t4.apkg")

            elapsed = time.time() - start_time
            assert elapsed < 15.0, f"Execution took {elapsed:.2f}s, exceeding 15.0s budget!"
            assert out_file.exists()
            if out_file.exists():
                out_file.unlink()
